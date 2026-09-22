#!/usr/bin/env python3
"""
文化祭レジ（ベビーカステラ）のスプレッドシート土台を .xlsx として書き出す。

    python3 apps/festival-pos/template/build_template.py

Google ドライブへアップロードすると Google スプレッドシートに変換される。
Apps Script を貼り付けて「レジ管理 → シートを初期化 / 修復」を実行すると、
書式と集計（商品別・担当者別・時間帯別）が最終形になる。

数式は Excel と Google スプレッドシートの両方で動くものだけを使う
（QUERY と A1:A 形式の開いた範囲は Google 専用なので、ここでは使わない）。
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

TITLE_BG, TITLE_FG = "1F2937", "FFFFFF"
HEADER_BG, HEADER_FG = "374151", "FFFFFF"
EDITABLE = "FFF8C4"
LOCKED = "F3F4F6"
BAND = "E5E7EB"
NOTE = "6B7280"
YEN = "¥#,##0"
DATETIME = "yyyy-mm-dd hh:mm:ss"

REPO = "https://github.com/ORCA-IRUKA-222/skiheil-task"

# 売上シート。Code.gs の SALES_HEADERS と必ず一致させること
SALES_HEADERS = [
    "受信日時", "取引ID", "明細番号", "端末ID", "担当者", "会計日時", "種別", "取消元取引ID",
    "商品名", "区分", "数量", "個数", "単価", "金額", "カタログVer",
    "取引合計", "お預かり", "お釣り", "検証", "備考",
]
ITEMS_HEADERS = ["商品名", "3個セット", "7個セット", "バラ1個", "販売状態", "備考"]
CLOSING_HEADERS = [
    "記録日時", "担当者", "端末ID", "開始時刻",
    "釣銭準備金", "現金売上(理論)", "理論現金", "実査現金", "差額", "取引件数", "備考",
]

THIN = Side(style="thin", color="D8DBE0")


def title_block(ws, title, description, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14, color=TITLE_FG)
    cell.fill = PatternFill("solid", fgColor=TITLE_BG)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    for col in range(1, width + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=TITLE_BG)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    note = ws.cell(row=2, column=1, value=description)
    note.font = Font(size=9, color=NOTE)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 34


def header_row(ws, row, headers):
    for i, name in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.font = Font(bold=True, color=HEADER_FG)
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def band(ws, row, width, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True)
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BAND)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- 使い方

def build_usage(wb):
    ws = wb.create_sheet("使い方")
    width = 8
    title_block(ws, "文化祭レジ（ベビーカステラ）｜使い方",
                "このスプレッドシートはレジアプリの土台です。Apps Script を貼り付けてデプロイすると、"
                "スマホやタブレットがレジになります。", width)

    lines = [
        ("h", "セットアップ（最初に1回だけ）"),
        ("n", "1", "「設定」シートのパスコードを必ず変更する。初期値 0000 のままにしないこと"),
        ("n", "2", "「商品」シートに味ごとの行を作り、3個セット・7個セット・バラ1個の価格を入れる"),
        ("n", "3", "拡張機能 → Apps Script を開き、5つのファイルを同名で貼り付ける"),
        ("i", "", "スクリプト: Code.gs / Setup.gs"),
        ("i", "", "HTML: index.html / style.html / script.html"),
        ("i", "", "マニフェスト: ⚙プロジェクトの設定 →「appsscript.json をエディタで表示する」にチェックして貼り付け"),
        ("n", "4", "スプレッドシートを開き直し、メニュー「レジ管理 → シートを初期化 / 修復」を実行する"),
        ("i", "", "ダッシュボードに商品別・担当者別・時間帯別の集計が追加されます"),
        ("n", "5", "メニュー「レジ管理 → 設定をチェック」で価格の妥当性を確認する"),
        ("n", "6", "デプロイ → 新しいデプロイ → 種類「ウェブアプリ」"),
        ("i", "", "次のユーザーとして実行: 自分／アクセスできるユーザー: 全員"),
        ("n", "7", "発行された URL をレジ端末で開く"),
        ("", "", ""),
        ("h", "料金ルール"),
        ("b", "", "3個ちょうど → その味の3個セット価格"),
        ("b", "", "7個ちょうど → その味の7個セット価格"),
        ("b", "", "それ以外 → 個数 × その味のバラ単価"),
        ("b", "", "味もセットも組み合わせられます（プレーン7個 ＋ チョコ3個×2 など）"),
        ("b", "", "バラ単価によっては6個が7個セットより高くなります。「設定をチェック」が味ごとに警告します"),
        ("", "", ""),
        ("h", "当日の運用"),
        ("b", "", "アプリを開いて担当者名・パスコード・釣銭準備金を入力する"),
        ("b", "", "電波が切れてもレジは動きますが、タブは閉じないこと（閉じると圏外では開き直せません）"),
        ("b", "", "味や価格を変えるときは「商品」シートを直すだけ。最大15秒で全端末に反映されます"),
        ("b", "", "売り切れの味は「販売状態」を停止中にするとレジからボタンが消えます"),
        ("b", "", "打ち間違いはアプリの「会計履歴・取消」から取り消します"),
        ("b", "", "閉店時はアプリの「レジ締め」で実際の現金と照合します"),
        ("", "", ""),
        ("h", "注意"),
        ("b", "", "「売上」シートは 1行 = 1明細です。行の編集・削除はしないでください（端末側と食い違います）"),
        ("b", "", "WebアプリURLを SNS などに貼らないでください"),
        ("b", "", "過去の売上金額は会計時の単価で固定されます。あとから価格を変えても動きません"),
        ("", "", ""),
        ("h", "ソースコードと詳しい説明"),
        ("b", "", REPO + "  の apps/festival-pos/（ブランチ: feature/festival-pos）"),
    ]

    row = 4
    for entry in lines:
        kind = entry[0]
        if kind == "h":
            band(ws, row, width, entry[1])
        elif kind in ("n", "i", "b"):
            num, text = entry[1], entry[2]
            if num:
                cell = ws.cell(row=row, column=1, value=num)
                cell.alignment = Alignment(horizontal="center")
                cell.font = Font(bold=True)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=width)
            cell = ws.cell(row=row, column=2, value=("・" + text if kind == "b" else text))
            cell.alignment = Alignment(wrap_text=True, vertical="center", indent=2 if kind == "i" else 0)
            if kind == "i":
                cell.font = Font(size=10, color=NOTE)
        row += 1

    widths(ws, {"A": 4, **{get_column_letter(c): 14 for c in range(2, width + 1)}})
    return ws


# ---------------------------------------------------------------- ダッシュボード

def build_dashboard(wb):
    ws = wb.create_sheet("ダッシュボード")
    width = 8
    title_block(ws, "ダッシュボード",
                "「売上」シートから自動集計されます。取消は数量・金額がマイナスで入るため、"
                "そのまま合計すれば純売上になります。", width)

    rng = lambda col: "'売上'!${0}$5:${0}$5000".format(col)
    cards = [
        ("売上金額", "=SUM({})".format(rng("N")), YEN),
        ("販売個数", "=SUM({})".format(rng("L")), '#,##0"個"'),
        ("会計件数", '=COUNTIFS({},1,{},"売上")'.format(rng("C"), rng("G")), '#,##0"件"'),
        ("取消件数", '=COUNTIFS({},1,{},"取消")'.format(rng("C"), rng("G")), '#,##0"件"'),
    ]
    for i, (label, formula, fmt) in enumerate(cards):
        col = i * 2 + 1
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        head = ws.cell(row=4, column=col, value=label)
        head.font = Font(bold=True, color=HEADER_FG)
        head.alignment = Alignment(horizontal="center")
        for c in (col, col + 1):
            ws.cell(row=4, column=c).fill = PatternFill("solid", fgColor=HEADER_BG)

        ws.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 1)
        value = ws.cell(row=5, column=col, value=formula)
        value.font = Font(bold=True, size=16)
        value.alignment = Alignment(horizontal="center", vertical="center")
        value.number_format = fmt
    ws.row_dimensions[5].height = 22
    ws.row_dimensions[6].height = 22

    band(ws, 8, width, "区分別")
    for i, name in enumerate(["区分", "数量", "売上金額"], start=1):
        cell = ws.cell(row=9, column=i, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=LOCKED)
    for i, label in enumerate(["3個セット", "7個セット", "バラ"]):
        row = 10 + i
        ws.cell(row=row, column=1, value=label)
        qty = ws.cell(row=row, column=2,
                      value='=SUMIFS({},{},"{}")'.format(rng("K"), rng("J"), label))
        qty.number_format = "#,##0"
        amount = ws.cell(row=row, column=3,
                         value='=SUMIFS({},{},"{}")'.format(rng("N"), rng("J"), label))
        amount.number_format = YEN

    band(ws, 14, width, "商品別・担当者別・時間帯別")
    ws.merge_cells(start_row=15, start_column=1, end_row=16, end_column=width)
    note = ws.cell(row=15, column=1,
                   value="メニュー「レジ管理 → シートを初期化 / 修復」を実行すると、"
                         "ここに商品別（味ごと）・担当者別・時間帯別の集計が追加されます。"
                         "Google スプレッドシート専用の QUERY 関数を使うため、このテンプレートには含めていません。")
    note.font = Font(size=10, color=NOTE)
    note.alignment = Alignment(wrap_text=True, vertical="top")

    widths(ws, {get_column_letter(c): 13 for c in range(1, width + 1)})
    return ws


# ---------------------------------------------------------------- 商品

def build_items(wb):
    ws = wb.create_sheet("商品")
    width = len(ITEMS_HEADERS)
    first, last = 5, 24

    title_block(ws, "商品",
                "味ごとに1行つくります。ここを直すと各端末のボタンと価格が自動で入れ替わります（最大15秒）。"
                "過去の売上金額は会計時の単価で固定されているので変わりません。"
                "「販売状態」を停止中にすると、その味のボタンがレジから消えます。", width)
    header_row(ws, 4, ITEMS_HEADERS)

    samples = [
        ["プレーン", 300, 600, 120, "販売中", ""],
        ["チョコ", 350, 700, 140, "販売中", ""],
    ]
    for r, values in enumerate(samples, start=first):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)

    for r in range(first, last + 1):
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=EDITABLE)
        for c in (2, 3, 4):
            cell = ws.cell(row=r, column=c)
            cell.fill = PatternFill("solid", fgColor=EDITABLE)
            cell.number_format = YEN
        ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=EDITABLE)

    dv = DataValidation(type="list", formula1='"販売中,停止中"', allow_blank=True)
    dv.error = "「販売中」か「停止中」を選んでください。"
    dv.errorTitle = "販売状態"
    ws.add_data_validation(dv)
    dv.add("E{}:E{}".format(first, last))

    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 3, end_column=width)
    note = ws.cell(row=last + 2, column=1,
                   value="※ 商品名がキーです。同じ名前を2行書くと、下の行は無視されます"
                         "（「レジ管理 → 設定をチェック」が警告します）。"
                         "※「セット以外は常にバラ単価」で計算します。"
                         "3個・7個ちょうどのときだけセット価格、それ以外は 個数 × バラ単価 です。")
    note.font = Font(size=10, color=NOTE)
    note.alignment = Alignment(wrap_text=True, vertical="top")

    widths(ws, {"A": 18, "B": 13, "C": 13, "D": 13, "E": 12, "F": 32})
    return ws


# ---------------------------------------------------------------- 設定

def build_settings(wb):
    ws = wb.create_sheet("設定")
    width = 4
    title_block(ws, "設定",
                "黄色セルだけを変更します。価格は「商品」シートで管理します。", width)
    header_row(ws, 4, ["キー（変更しない）", "設定項目", "値", "説明"])

    rows = [
        ("SHOP_NAME", "店名", "ベビーカステラ", "アプリ上部に表示", True),
        ("PASSCODE", "パスコード", "0000",
         "⚠ 必ず変更してください。テンプレート共通の初期値です", True),
        ("CATALOG_VERSION", "カタログバージョン", 1, "自動更新。手で編集しない", False),
        ("CATALOG_UPDATED", "最終更新日時", "", "自動更新。手で編集しない", False),
    ]
    for i, (key, label, value, desc, editable) in enumerate(rows):
        row = 5 + i
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=label)
        cell = ws.cell(row=row, column=3, value=value)
        cell.fill = PatternFill("solid", fgColor=EDITABLE if editable else LOCKED)
        if key == "PASSCODE":
            cell.number_format = "@"  # 先頭0が消えないよう文字列扱い
            cell.font = Font(bold=True, color="B91C1C")
        desc_cell = ws.cell(row=row, column=4, value=desc)
        if key == "PASSCODE":
            desc_cell.font = Font(bold=True, color="B91C1C")

    widths(ws, {"A": 20, "B": 18, "C": 22, "D": 46})
    return ws


# ---------------------------------------------------------------- 売上

def build_sales(wb):
    ws = wb.create_sheet("売上")
    width = len(SALES_HEADERS)
    title_block(ws, "売上",
                "レジ端末から自動で追記されます。1行 = 1明細（商品×区分）で、"
                "同じ取引IDの行がひとつの会計です。"
                "取引合計・お預かり・お釣り・検証は、その取引の先頭明細行にだけ入ります。"
                "行の編集・削除はしないでください（端末側と食い違います）。", width)
    header_row(ws, 4, SALES_HEADERS)

    # 空行の数値書式はあえて付けない。「レジ管理 → シートを初期化 / 修復」が
    # 列全体に付け直すため、ここで付けるとファイルが無駄に大きくなるだけ。

    widths(ws, {"A": 19, "B": 30, "C": 9, "D": 12, "E": 12, "F": 19, "G": 8, "H": 30,
                "I": 16, "J": 12, "K": 8, "L": 8, "M": 10, "N": 11, "O": 11,
                "P": 11, "Q": 11, "R": 11, "S": 22, "T": 20})
    return ws


# ---------------------------------------------------------------- レジ締め

def build_closing(wb):
    ws = wb.create_sheet("レジ締め")
    width = len(CLOSING_HEADERS)
    title_block(ws, "レジ締め",
                "アプリの「レジ締め」から記録されます。"
                "理論現金 = 釣銭準備金 + そのシフトの現金売上。差額がマイナスなら現金が足りません。", width)
    header_row(ws, 4, CLOSING_HEADERS)

    widths(ws, {"A": 19, "B": 12, "C": 12, "D": 19, "E": 13, "F": 15,
                "G": 12, "H": 12, "I": 11, "J": 10, "K": 24})
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    build_usage(wb)
    build_dashboard(wb)
    build_items(wb)
    build_settings(wb)
    build_sales(wb)
    build_closing(wb)
    wb.active = 0

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "文化祭レジ_ベビーカステラ_テンプレート.xlsx")
    wb.save(out)
    print("{} ({:,} bytes)".format(out, os.path.getsize(out)))
    print("シート: " + " / ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
